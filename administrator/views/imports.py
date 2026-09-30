import logging

from django.contrib import messages
from django.contrib.auth.models import User
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, render

from student.models import StudentProfile
from teacher.models import Module, ModuleAssignment, ModuleEnrollment, TeacherProfile

from ..models import AdminProfile
from ._auth import admin_required

logger = logging.getLogger(__name__)

STUDENT_IMPORT_HEADER_ALIASES = {
    "prenom": "first_name",
    "prénom": "first_name",
    "first_name": "first_name",
    "nom": "last_name",
    "last_name": "last_name",
    "email": "email",
    "e-mail": "email",
    "id etudiant": "student_id",
    "id étudiant": "student_id",
    "student_id": "student_id",
    "cne": "student_id",
    "annee": "year_of_study",
    "année": "year_of_study",
    "year": "year_of_study",
    "year_of_study": "year_of_study",
    "departement": "department",
    "département": "department",
    "department": "department",
}

STUDENT_IMPORT_REQUIRED_FIELDS = [
    "first_name",
    "last_name",
    "email",
    "student_id",
    "year_of_study",
    "department",
]

TEACHER_IMPORT_HEADER_ALIASES = {
    "prenom": "first_name",
    "prénom": "first_name",
    "first_name": "first_name",
    "nom": "last_name",
    "last_name": "last_name",
    "email": "email",
    "e-mail": "email",
    "id enseignant": "teacher_id",
    "teacher_id": "teacher_id",
    "id": "teacher_id",
    "departement": "department",
    "département": "department",
    "department": "department",
}

TEACHER_IMPORT_REQUIRED_FIELDS = ["first_name", "last_name", "email", "department"]


def _generate_unique_username(email):
    """Build a unique username from the email local part"""
    base = email.split("@")[0].lower().replace(" ", ".")
    username = base
    suffix = 1
    while User.objects.filter(username=username).exists():
        suffix += 1
        username = f"{base}{suffix}"
    return username


@admin_required
def import_students(request):
    """Bulk-create student accounts from an Excel (.xlsx) roster"""
    import openpyxl
    from django.db import transaction
    from django.utils.crypto import get_random_string

    admin = get_object_or_404(AdminProfile, user=request.user)
    modules = Module.objects.filter(is_active=True).order_by("code")

    created = []
    skipped = []
    errors = []
    import_done = False

    if request.method == "POST":
        excel_file = request.FILES.get("excel_file")
        module_id = request.POST.get("module_id", "")

        module = None
        if module_id:
            module = get_object_or_404(Module, id=module_id)

        if not excel_file:
            messages.error(request, "Veuillez sélectionner un fichier Excel (.xlsx).")
        elif not excel_file.name.lower().endswith(".xlsx"):
            messages.error(
                request,
                "Format non supporté. Utilisez un fichier .xlsx (modèle téléchargeable ci-dessous).",
            )
        else:
            try:
                workbook = openpyxl.load_workbook(excel_file, read_only=True, data_only=True)
                sheet = workbook.active

                rows = sheet.iter_rows(values_only=True)
                header_row = next(rows, None)
                if header_row is None:
                    messages.error(request, "Le fichier est vide.")
                    raise StopIteration

                # Map column index -> field name using header aliases
                column_map = {}
                for index, header in enumerate(header_row):
                    if header is None:
                        continue
                    key = str(header).strip().lower()
                    if key in STUDENT_IMPORT_HEADER_ALIASES:
                        column_map[index] = STUDENT_IMPORT_HEADER_ALIASES[key]

                missing_columns = [
                    f for f in STUDENT_IMPORT_REQUIRED_FIELDS if f not in column_map.values()
                ]
                if missing_columns:
                    messages.error(
                        request,
                        "Colonnes manquantes dans le fichier: Prénom, Nom, Email, ID Étudiant, Année, Département sont requises. "
                        "Téléchargez le modèle pour voir le format attendu.",
                    )
                else:
                    for row_number, row in enumerate(rows, start=2):
                        # Skip fully empty rows
                        if row is None or all(
                            cell is None or str(cell).strip() == "" for cell in row
                        ):
                            continue

                        data = {}
                        for index, field in column_map.items():
                            value = row[index] if index < len(row) else None
                            data[field] = str(value).strip() if value is not None else ""

                        # Validate required values
                        missing = [f for f in STUDENT_IMPORT_REQUIRED_FIELDS if not data.get(f)]
                        if missing:
                            errors.append(
                                {
                                    "row": row_number,
                                    "reason": "Valeurs manquantes (toutes les colonnes sont obligatoires)",
                                }
                            )
                            continue

                        # Year must be 3, 4 or 5
                        try:
                            year = int(float(data["year_of_study"]))
                        except ValueError:
                            year = None
                        if year not in [3, 4, 5]:
                            errors.append(
                                {
                                    "row": row_number,
                                    "reason": f"Année invalide '{data['year_of_study']}' (3, 4 ou 5 attendu)",
                                }
                            )
                            continue

                        email = data["email"].lower()
                        if User.objects.filter(email__iexact=email).exists():
                            skipped.append(
                                {
                                    "row": row_number,
                                    "identity": email,
                                    "reason": "Email déjà utilisé",
                                }
                            )
                            continue
                        if StudentProfile.objects.filter(student_id=data["student_id"]).exists():
                            skipped.append(
                                {
                                    "row": row_number,
                                    "identity": data["student_id"],
                                    "reason": "ID étudiant déjà utilisé",
                                }
                            )
                            continue

                        password = get_random_string(10)
                        try:
                            with transaction.atomic():
                                user = User.objects.create_user(
                                    username=_generate_unique_username(email),
                                    email=email,
                                    password=password,
                                    first_name=data["first_name"],
                                    last_name=data["last_name"],
                                    is_active=True,
                                )
                                student = StudentProfile.objects.create(
                                    user=user,
                                    student_id=data["student_id"],
                                    year_of_study=year,
                                    department=data["department"],
                                )
                                if module:
                                    ModuleEnrollment.objects.create(
                                        student=student, module=module, is_active=True
                                    ).attach_published_assignments()
                        except Exception as e:
                            errors.append(
                                {
                                    "row": row_number,
                                    "reason": f"Erreur lors de la création: {str(e)}",
                                }
                            )
                            continue

                        created.append(
                            {
                                "name": f"{data['first_name']} {data['last_name']}",
                                "student_id": data["student_id"],
                                "email": email,
                                "username": user.username,
                                "password": password,
                            }
                        )

                    import_done = True
                    if created:
                        enrolled_note = f" et inscrits au module {module.code}" if module else ""
                        messages.success(
                            request,
                            f"{len(created)} étudiant(s) créé(s){enrolled_note}. "
                            "Notez les mots de passe ci-dessous: ils ne seront plus affichés.",
                        )
                    if skipped:
                        messages.warning(
                            request, f"{len(skipped)} ligne(s) ignorée(s) (comptes déjà existants)."
                        )
                    if errors:
                        messages.error(request, f"{len(errors)} ligne(s) en erreur.")
                    if not created and not skipped and not errors:
                        messages.info(request, "Aucune ligne de données trouvée dans le fichier.")
            except StopIteration:
                pass
            except Exception:
                logger.exception("Impossible de lire le fichier Excel")
                messages.error(
                    request,
                    "Impossible de lire le fichier Excel. Réessayez ; si le problème persiste, contactez l'administration.",
                )

    context = {
        "admin": admin,
        "modules": modules,
        "created": created,
        "skipped": skipped,
        "errors": errors,
        "import_done": import_done,
    }

    return render(request, "administrator/import_students.html", context)


@admin_required
def download_student_import_template(request):
    """Downloadable Excel template for the student import"""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Étudiants"
    sheet.append(["Prénom", "Nom", "Email", "ID Étudiant", "Année", "Département"])
    sheet.append(["Amina", "Benali", "amina.benali@example.com", "E12345", 3, "Informatique"])
    sheet.append(
        ["Youssef", "El Amrani", "youssef.elamrani@example.com", "E12346", 4, "Génie Civil"]
    )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="modele_import_etudiants.xlsx"'
    workbook.save(response)
    return response


@admin_required
def import_teachers(request):
    """Bulk-create teacher accounts from an Excel (.xlsx) roster"""
    import openpyxl
    from django.db import transaction
    from django.utils.crypto import get_random_string

    admin = get_object_or_404(AdminProfile, user=request.user)
    modules = Module.objects.filter(is_active=True).order_by("code")

    created = []
    skipped = []
    errors = []
    import_done = False

    if request.method == "POST":
        excel_file = request.FILES.get("excel_file")
        module_id = request.POST.get("module_id", "")

        module = None
        if module_id:
            module = get_object_or_404(Module, id=module_id)

        if not excel_file:
            messages.error(request, "Veuillez sélectionner un fichier Excel (.xlsx).")
        elif not excel_file.name.lower().endswith(".xlsx"):
            messages.error(
                request,
                "Format non supporté. Utilisez un fichier .xlsx (modèle téléchargeable ci-dessous).",
            )
        else:
            try:
                workbook = openpyxl.load_workbook(excel_file, read_only=True, data_only=True)
                sheet = workbook.active

                rows = sheet.iter_rows(values_only=True)
                header_row = next(rows, None)
                if header_row is None:
                    messages.error(request, "Le fichier est vide.")
                    raise StopIteration

                column_map = {}
                for index, header in enumerate(header_row):
                    if header is None:
                        continue
                    key = str(header).strip().lower()
                    if key in TEACHER_IMPORT_HEADER_ALIASES:
                        column_map[index] = TEACHER_IMPORT_HEADER_ALIASES[key]

                missing_columns = [
                    f for f in TEACHER_IMPORT_REQUIRED_FIELDS if f not in column_map.values()
                ]
                if missing_columns:
                    messages.error(
                        request,
                        "Colonnes manquantes dans le fichier: Prénom, Nom, Email, Département sont requises "
                        "(ID Enseignant optionnel). Téléchargez le modèle pour voir le format attendu.",
                    )
                else:
                    for row_number, row in enumerate(rows, start=2):
                        if row is None or all(
                            cell is None or str(cell).strip() == "" for cell in row
                        ):
                            continue

                        data = {}
                        for index, field in column_map.items():
                            value = row[index] if index < len(row) else None
                            data[field] = str(value).strip() if value is not None else ""

                        missing = [f for f in TEACHER_IMPORT_REQUIRED_FIELDS if not data.get(f)]
                        if missing:
                            errors.append(
                                {
                                    "row": row_number,
                                    "reason": "Valeurs manquantes (Prénom, Nom, Email et Département sont obligatoires)",
                                }
                            )
                            continue

                        email = data["email"].lower()
                        teacher_id = data.get("teacher_id", "")

                        if User.objects.filter(email__iexact=email).exists():
                            skipped.append(
                                {
                                    "row": row_number,
                                    "identity": email,
                                    "reason": "Email déjà utilisé",
                                }
                            )
                            continue
                        if (
                            teacher_id
                            and TeacherProfile.objects.filter(teacher_id=teacher_id).exists()
                        ):
                            skipped.append(
                                {
                                    "row": row_number,
                                    "identity": teacher_id,
                                    "reason": "ID enseignant déjà utilisé",
                                }
                            )
                            continue

                        password = get_random_string(10)
                        try:
                            with transaction.atomic():
                                user = User.objects.create_user(
                                    username=_generate_unique_username(email),
                                    email=email,
                                    password=password,
                                    first_name=data["first_name"],
                                    last_name=data["last_name"],
                                    is_active=True,
                                )
                                teacher = TeacherProfile.objects.create(
                                    user=user,
                                    teacher_id=teacher_id or None,
                                    department=data["department"],
                                )
                                if module:
                                    ModuleAssignment.objects.create(
                                        teacher=teacher,
                                        module=module,
                                        assigned_by=request.user,
                                        is_active=True,
                                    )
                        except Exception as e:
                            errors.append(
                                {
                                    "row": row_number,
                                    "reason": f"Erreur lors de la création: {str(e)}",
                                }
                            )
                            continue

                        created.append(
                            {
                                "name": f"{data['first_name']} {data['last_name']}",
                                "teacher_id": teacher_id or "—",
                                "email": email,
                                "username": user.username,
                                "password": password,
                            }
                        )

                    import_done = True
                    if created:
                        assigned_note = f" et assignés au module {module.code}" if module else ""
                        messages.success(
                            request,
                            f"{len(created)} enseignant(s) créé(s){assigned_note}. "
                            "Notez les mots de passe ci-dessous: ils ne seront plus affichés.",
                        )
                    if skipped:
                        messages.warning(
                            request, f"{len(skipped)} ligne(s) ignorée(s) (comptes déjà existants)."
                        )
                    if errors:
                        messages.error(request, f"{len(errors)} ligne(s) en erreur.")
                    if not created and not skipped and not errors:
                        messages.info(request, "Aucune ligne de données trouvée dans le fichier.")
            except StopIteration:
                pass
            except Exception:
                logger.exception("Impossible de lire le fichier Excel")
                messages.error(
                    request,
                    "Impossible de lire le fichier Excel. Réessayez ; si le problème persiste, contactez l'administration.",
                )

    context = {
        "admin": admin,
        "modules": modules,
        "created": created,
        "skipped": skipped,
        "errors": errors,
        "import_done": import_done,
    }

    return render(request, "administrator/import_teachers.html", context)


@admin_required
def download_teacher_import_template(request):
    """Downloadable Excel template for the teacher import"""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Enseignants"
    sheet.append(["Prénom", "Nom", "Email", "ID Enseignant", "Département"])
    sheet.append(["Karim", "Alaoui", "karim.alaoui@example.com", "T1001", "Informatique"])
    sheet.append(["Salma", "Bennis", "salma.bennis@example.com", "", "Génie Civil"])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="modele_import_enseignants.xlsx"'
    workbook.save(response)
    return response
